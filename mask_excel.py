"""
mask_excel.py
=============
Excel (.xlsx) dosyaları üzerinde, bir eşleştirme tablosundaki
(anahtar_kelimeler.csv) anahtar kelimeleri karşılıklarıyla değiştirir.
mask_pdf.py ile AYNI mantık ve AYNI eşleştirme dosyası formatını kullanır;
tek fark, PDF yerine Excel üzerinde çalışması ve bir Excel dosyasındaki
TÜM SAYFALARI (tek sayfa değil) taramasıdır.

NASIL ÇALIŞIR
-------------
1. Dosyadaki HER SAYFA taranır (tek sayfayla sınırlı değil) -- hücre
   değerleri, sayfa adları, hücre yorumları (comment) ve sayfa üstbilgi/
   altbilgileri (varsa) dahil.
2. Bir hücre anahtar kelimeyle eşleşirse, hücrenin değeri GERÇEKTEN
   karşılığıyla DEĞİŞTİRİLİR (PDF'teki "üzerini kapatma" derdi Excel'de
   yok -- bir hücrenin değerini değiştirmek, eskisini olduğu gibi siler).
3. Kaydettikten sonra, dosya HAM DÜZEYDE (xlsx aslında bir ZIP arşividir;
   içindeki tüm XML parçaları -- paylaşılan metinler, grafik başlıkları,
   belge özellikleri vb.) taranır. openpyxl'in hücre bazlı geçişinin
   atlayabileceği bir yer (örn. bir grafik başlığı) varsa, doğrudan o XML
   parçası üzerinde metin düzeyinde onarılır.
4. Script, kaydettiği HER dosyayı otomatik olarak yeniden açıp orijinal
   anahtar kelimelerin dosyanın hiçbir yerinde (hiçbir sayfada, hiçbir iç
   XML parçasında) kalmadığını doğrular ve sonucu raporlar.

anahtar_kelimeler.csv dosyası mask_pdf.py ile AYNI dosyadır -- aynı
eşleştirmeyi (örn. "Kadıköy Bağdat Caddesi Şubesi" -> "SUBE_001") hem PDF
hem Excel'de kullanmak, iki dosya arasında tutarlılığı sağlar. İsterseniz
AYARLAR bölümünden farklı bir CSV dosyası da gösterebilirsiniz.

KULLANIM
--------
Komut satırı argümanı YOK. Aşağıdaki "AYARLAR" bölümündeki yolları bu
dosyanın içinde düzenleyip kaydedin, sonra çalıştırın:

    python mask_excel.py

GEREKSİNİMLER
-------------
    pip install openpyxl

BİLİNEN KISITLAR
-----------------
- Yalnızca .xlsx (ve .xlsm) desteklenir. Eski .xls formatındaki dosyaları
  önce Excel'de "Farklı Kaydet" ile .xlsx'e dönüştürün.
- Anahtar kelimeniz bir hücre referansına benziyorsa (örn. "A1"), formül
  içeren hücrelerde yanlış eşleşme riski teorik olarak vardır -- gerçek
  şube adı/kodu gibi anahtar kelimelerde bu risk yok denecek kadar azdır.
"""
import csv
import sys
import zipfile
from pathlib import Path

import openpyxl

# =============================================================================
# AYARLAR — çalıştırmadan önce bu yolları kendi durumunuza göre güncelleyin
# =============================================================================
# EXCEL_INPUT_PATH: Tek bir .xlsx dosyası ya da .xlsx dosyalarının bulunduğu
#                    bir klasör olabilir (klasörse içindeki tüm .xlsx işlenir).
EXCEL_INPUT_PATH = "istatistikler/"

# ANAHTAR_KELIMELER_CSV: mask_pdf.py ile AYNI dosya (aynı klasördeyse aynı
#                         eşleştirme otomatik kullanılır). Farklı bir
#                         eşleştirme istiyorsanız burada başka bir yol verin.
ANAHTAR_KELIMELER_CSV = "anahtar_kelimeler.csv"

# OUTPUT_DIR: Maskelenmiş Excel dosyalarının yazılacağı klasör (yoksa oluşturulur).
OUTPUT_DIR = "maskeli/"

# DRY_RUN: True yapılırsa dosyalar DEĞİŞTİRİLMEZ, yalnızca kaç eşleşme
#          bulunduğu konsolda gösterilir.
DRY_RUN = False
# =============================================================================

# Türkçe'de İngilizce'den farklı olarak 4 ayrı "i" harfi vardır: İ/i (noktalı)
# ve I/ı (noktasız). Python'un standart .upper()/.lower()'ı bu ayrımı doğru
# yapmaz (örn. "İstanbul".lower() -> "i̇stanbul" gibi hatalı/birleşik
# karakterli bir sonuç verir, "istanbul" değil). Bu yüzden aramayı, kelimenin
# gerçekçi tüm yazım/büyük-küçük harf varyantlarını deneyerek yapıyoruz
# (mask_pdf.py ile birebir aynı mantık).
_TR_LOWER_MAP = str.maketrans({"İ": "i", "I": "ı"})
_TR_UPPER_MAP = str.maketrans({"i": "İ", "ı": "I"})
_ASCII_FOLD_MAP = str.maketrans({"İ": "I", "ı": "i"})  # Türkçe klavyesiz/ASCII yazım


def _tr_lower(s: str) -> str:
    return s.translate(_TR_LOWER_MAP).lower()


def _tr_upper(s: str) -> str:
    return s.translate(_TR_UPPER_MAP).upper()


def case_variants(keyword: str) -> list[str]:
    """Bir anahtar kelimenin Excel'de karşılaşılabilecek gerçekçi tüm
    büyük/küçük harf ve Türkçe İ/I/ı/i yazım varyantlarını üretir."""
    ascii_folded = keyword.translate(_ASCII_FOLD_MAP)
    variants = {
        keyword,
        keyword.upper(), keyword.lower(), keyword.title(),
        _tr_upper(keyword), _tr_lower(keyword),
        ascii_folded, ascii_folded.upper(), ascii_folded.lower(), ascii_folded.title(),
    }
    return [v for v in variants if v]


def load_mapping(csv_path: Path) -> list[tuple[str, str]]:
    """CSV'den (anahtar_kelime, karsilik) çiftlerini okur.
    Boş/placeholder anahtar kelimeler (henüz doldurulmamış satırlar) atlanır."""
    pairs = []
    with open(csv_path, encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            kw = (row.get("anahtar_kelime") or "").strip()
            val = (row.get("karsilik") or "").strip()
            if not kw or kw.startswith("ANAHTAR_KELIME_"):
                continue  # henüz doldurulmamış placeholder satır
            pairs.append((kw, val))
    # Uzun anahtar kelimeleri önce işle: biri diğerinin alt dizesiyse
    # (örn. "Ankara" ve "Ankara Kızılay Şubesi") uzun olanın önce
    # eşleşmesi yanlış kısmi değişimi önler.
    pairs.sort(key=lambda p: len(p[0]), reverse=True)
    return pairs


def mask_text_value(text: str, mapping: list[tuple[str, str]], value_transform=None) -> tuple[str, dict]:
    """Bir metin içindeki tüm anahtar kelime geçişlerini (tüm Türkçe
    büyük/küçük harf varyantlarıyla) karşılıklarıyla değiştirir.
    (maskelenmiş_metin, {anahtar_kelime: eşleşme_sayısı}) döner.

    ÖNEMLİ: TEK GEÇİŞTE (soldan sağa) tarar -- sırayla text.replace()
    çağırmak YERİNE. Sebep: bir karşılık değeri (örn. "İstanbul" ->
    "SUBE_ISTANBUL") anahtar kelimenin kendisini ("ISTANBUL") alt dize
    olarak içerebilir. Sıralı replace() kullanılsaydı, "Istanbul" önce
    "SUBE_ISTANBUL" ile değiştirilir, sonra AYNI DÖNGÜ İÇİNDE "ISTANBUL"
    varyantı bu YENİ metnin içinde ("SUBE_ISTANBUL" içindeki "ISTANBUL")
    tekrar bulunup "SUBE_SUBE_ISTANBUL" gibi ÇİFTE MASKELEMEYE yol açardı
    (bizzat test ederek bu hatayı bulduk). Tek geçişli tarama, metindeki
    her pozisyonu yalnızca bir kez işlediği ve zaten yazılmış karşılık
    metnini asla geriye dönüp yeniden taramadığı için bu sorunu tamamen
    ortadan kaldırır."""
    all_variants = []
    for keyword, value in mapping:
        out_value = value_transform(value) if value_transform else value
        for variant in case_variants(keyword):
            if variant:
                all_variants.append((variant, out_value, keyword))
    # En uzun varyant önce denenir (bir varyant başka birinin alt dizesiyse
    # yanlış kısmi eşleşmeyi önlemek için).
    all_variants.sort(key=lambda v: len(v[0]), reverse=True)

    counts: dict[str, int] = {}
    result = []
    i, n = 0, len(text)
    while i < n:
        matched = False
        for variant, out_value, keyword in all_variants:
            if text.startswith(variant, i):
                result.append(out_value)
                counts[keyword] = counts.get(keyword, 0) + 1
                i += len(variant)
                matched = True
                break
        if not matched:
            result.append(text[i])
            i += 1
    return "".join(result), counts


def mask_cell_value(value, mapping: list[tuple[str, str]]):
    """Bir hücre değerini (metin veya sayı olabilir) maskeler.
    (yeni_değer, {anahtar_kelime: eşleşme_sayısı}) döner; eşleşme yoksa
    değer değişmeden döner."""
    if value is None:
        return value, {}
    if isinstance(value, str):
        return mask_text_value(value, mapping)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Örn. şube kodu Excel'de metin değil SAYI olarak girilmiş olabilir
        # (47582931 gibi). Sayıyı metne çevirip aynı mantıkla arıyoruz.
        text = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
        masked_text, counts = mask_text_value(text, mapping)
        if counts:
            return masked_text, counts  # karşılık genelde alfanumerik olduğu için hücre artık metne döner
        return value, {}
    return value, {}


def mask_workbook(wb: "openpyxl.Workbook", mapping: list[tuple[str, str]]) -> dict:
    """Çalışma kitabındaki TÜM SAYFALARI (sayfa adı, hücreler, hücre
    yorumları, üstbilgi/altbilgi dahil) tarayıp yerinde maskeler.
    {anahtar_kelime: toplam_eşleşme_sayısı} döner."""
    total_counts: dict[str, int] = {kw: 0 for kw, _ in mapping}

    def add(counts: dict):
        for kw, c in counts.items():
            total_counts[kw] = total_counts.get(kw, 0) + c

    for ws in wb.worksheets:
        new_title, counts = mask_text_value(ws.title, mapping)
        if counts:
            ws.title = new_title[:31]  # Excel sayfa adı sınırı: 31 karakter
            add(counts)

        for row in ws.iter_rows():
            for cell in row:
                new_value, counts = mask_cell_value(cell.value, mapping)
                if counts:
                    cell.value = new_value
                    add(counts)
                if cell.comment is not None and cell.comment.text:
                    new_text, c_counts = mask_text_value(cell.comment.text, mapping)
                    if c_counts:
                        cell.comment.text = new_text
                        add(c_counts)

        for hf_attr in ("oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader", "firstFooter"):
            hf = getattr(ws, hf_attr, None)
            if hf is None:
                continue
            for part_attr in ("left", "center", "right"):
                part = getattr(hf, part_attr, None)
                if part is None or not part.text:
                    continue
                new_text, counts = mask_text_value(part.text, mapping)
                if counts:
                    part.text = new_text
                    add(counts)

    return total_counts


def scan_zip_for_keywords(path: Path, mapping: list[tuple[str, str]]) -> dict[str, list[str]]:
    """xlsx dosyası aslında bir ZIP arşividir. Bu, içindeki HER XML
    parçasını (paylaşılan metinler, grafikler, belge özellikleri, ilişki
    dosyaları vb.) ham metin düzeyinde tarayıp, hangi anahtar kelimenin
    hangi iç dosyada hâlâ göründüğünü döndürür (openpyxl'in hücre bazlı
    geçişinin atlayabileceği yerleri -- örn. bir grafik başlığı -- yakalamak
    içindir). {anahtar_kelime: [iç_dosya_adı, ...]} döner; boşsa temiz.

    ÖNEMLİ (yanlış pozitif düzeltmesi): karşılık (value) değerleri anahtar
    kelimeyi alt dize olarak içerebilir (örn. "İstanbul" -> "SUBE_ISTANBUL"
    karşılığının içinde "ISTANBUL" geçer). Bu durumda, doğru şekilde
    maskelenmiş metin yanlışlıkla "hâlâ sızıyor" gibi görünür. Bunu önlemek
    için, taramadan ÖNCE içerikten zaten uygulanmış TÜM karşılık
    değerlerini çıkarıyoruz -- geriye kalanda kelime hâlâ varsa, bu gerçek
    bir sızıntıdır."""
    hits: dict[str, list[str]] = {}
    with zipfile.ZipFile(path, "r") as zf:
        parts = {}
        for name in zf.namelist():
            if not (name.endswith(".xml") or name.endswith(".rels")):
                continue
            try:
                parts[name] = zf.read(name).decode("utf-8", errors="ignore")
            except Exception:
                continue

    replacement_values = [value for _, value in mapping if value]

    for keyword, _ in mapping:
        for variant in case_variants(keyword):
            for name, content in parts.items():
                scrubbed = content
                for value in replacement_values:
                    scrubbed = scrubbed.replace(value, "")
                if variant in scrubbed:
                    hits.setdefault(keyword, [])
                    if name not in hits[keyword]:
                        hits[keyword].append(name)
    return hits


def repair_zip_leaks(path: Path, mapping: list[tuple[str, str]], leaked: dict[str, list[str]]) -> None:
    """scan_zip_for_keywords'ün bulduğu, openpyxl'in hücre bazlı geçişinin
    değiştiremediği iç XML parçalarını (örn. bir grafik başlığı, belge
    özellikleri) doğrudan metin düzeyinde onarır -- xlsx'i ZIP olarak
    yeniden yazarak.

    ÖNEMLİ (çifte maskelemeyi önleme): bir karşılık değeri anahtar kelimeyi
    alt dize olarak içerebilir (örn. "SUBE_ISTANBUL" içinde "ISTANBUL"
    geçer). Zaten doğru şekilde maskelenmiş bir metni tekrar taramak
    "SUBE_SUBE_ISTANBUL" gibi bir çifte maskelemeye yol açar. Bunu önlemek
    için, önce içerikteki tüm karşılık değerlerini geçici (gerçek metinde
    asla geçmeyecek) token'larla koruyup, taramayı öyle yapıyor, sonra
    token'ları orijinal karşılık değerlerine geri çeviriyoruz."""
    affected_files = {name for names in leaked.values() for name in names}
    if not affected_files:
        return

    with zipfile.ZipFile(path, "r") as zf:
        items = {name: zf.read(name) for name in zf.namelist()}

    replacement_values = [value for _, value in mapping if value]

    def xml_escape(v: str) -> str:
        return v.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    for name in affected_files:
        try:
            content = items[name].decode("utf-8")
        except UnicodeDecodeError:
            continue

        # 1) Bu dosyada ZATEN doğru şekilde uygulanmış karşılık değerlerini
        #    (örn. openpyxl'in daha önce yazdığı "SUBE_ISTANBUL") geçici,
        #    gerçek metinde asla geçmeyecek token'larla koru -- aksi halde
        #    tek geçişli tarama bile bunların İÇİNDEKİ "ISTANBUL" alt
        #    dizesini yeni bir eşleşme sanıp tekrar maskeler.
        placeholders: dict[str, str] = {}
        protected = content
        for idx, value in enumerate(replacement_values):
            if value and value in protected:
                token = f"\x00PH_{idx}\x00"
                placeholders[token] = value
                protected = protected.replace(value, token)

        # 2) Geriye kalan (korunmamış = gerçekten maskelenmemiş) metinde,
        #    aynı tek-geçişli/güvenli algoritmayla ara ve değiştir.
        protected, _ = mask_text_value(protected, mapping, value_transform=xml_escape)

        # 3) Token'ları orijinal (XML-kaçışlı) karşılık değerlerine geri çevir.
        for token, value in placeholders.items():
            protected = protected.replace(token, xml_escape(value))

        items[name] = protected.encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in items.items():
            zf.writestr(name, data)


def mask_excel(input_path: Path, output_path: Path, mapping: list[tuple[str, str]], dry_run: bool = False) -> dict:
    wb = openpyxl.load_workbook(input_path)
    n_sheets = len(wb.worksheets)
    counts = mask_workbook(wb, mapping)  # bellekteki kopyayı değiştirir; orijinal dosyaya dokunmaz

    if not dry_run:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)

    counts["_n_sheets"] = n_sheets  # main()'de bilgi amaçlı kullanılır, anahtar kelime değil
    return counts


def main():
    mapping_path = Path(ANAHTAR_KELIMELER_CSV)
    if not mapping_path.exists():
        print(f"[hata] Eşleştirme dosyası bulunamadı: {mapping_path}", file=sys.stderr)
        print("       AYARLAR bölümündeki ANAHTAR_KELIMELER_CSV yolunu kontrol edin.", file=sys.stderr)
        sys.exit(1)
    mapping = load_mapping(mapping_path)
    if not mapping:
        print(
            f"[hata] {mapping_path} içinde doldurulmuş satır yok. "
            f"'anahtar_kelime' sütununu ANAHTAR_KELIME_XX placeholder'ları yerine "
            f"gerçek kelimelerle doldurun.",
            file=sys.stderr,
        )
        sys.exit(1)
    print(f"[bilgi] {len(mapping)} anahtar kelime/karşılık çifti yüklendi.")

    in_path = Path(EXCEL_INPUT_PATH)
    if not in_path.exists():
        print(f"[hata] {in_path} bulunamadı. AYARLAR bölümündeki EXCEL_INPUT_PATH yolunu kontrol edin.", file=sys.stderr)
        sys.exit(1)
    files = sorted(in_path.glob("*.xlsx")) if in_path.is_dir() else [in_path]

    xls_files = sorted(in_path.glob("*.xls")) if in_path.is_dir() else ([in_path] if in_path.suffix == ".xls" else [])
    if xls_files:
        print(
            f"[uyarı] {len(xls_files)} adet eski formatta (.xls) dosya bulundu, bunlar DESTEKLENMİYOR: "
            f"{[f.name for f in xls_files]} -- önce Excel'de 'Farklı Kaydet' ile .xlsx'e dönüştürün.",
            file=sys.stderr,
        )

    if not files:
        print(f"[hata] {in_path} altında .xlsx bulunamadı", file=sys.stderr)
        sys.exit(1)

    out_dir = Path(OUTPUT_DIR)
    total_counts = {kw: 0 for kw, _ in mapping}
    all_leak_warnings = []
    original_keywords = [kw for kw, _ in mapping]

    for f in files:
        out_path = out_dir / f"maskeli_{f.name}"
        counts = mask_excel(f, out_path, mapping, dry_run=DRY_RUN)
        n_sheets = counts.pop("_n_sheets")
        found = sum(counts.values())
        action = "bulundu (dry-run, dosya değiştirilmedi)" if DRY_RUN else f"değiştirildi -> {out_path.name}"
        print(f"[{f.name}] {n_sheets} sayfa tarandı, {found} eşleşme {action}")
        for kw, c in counts.items():
            total_counts[kw] += c

        if not DRY_RUN:
            leaked = scan_zip_for_keywords(out_path, mapping)
            if leaked:
                repair_zip_leaks(out_path, mapping, leaked)
                leaked = scan_zip_for_keywords(out_path, mapping)
            if leaked:
                details = "; ".join(f"{kw} ({', '.join(files_)})" for kw, files_ in leaked.items())
                all_leak_warnings.append(f"[{f.name}] HÂLÂ DOSYADA BULUNAN: {details}")
            else:
                print(f"  [doğrulandı] {out_path.name}: orijinal anahtar kelimelerin hiçbiri dosyada kalmadı.")

    print("\n[özet] Anahtar kelime başına toplam eşleşme sayısı:")
    for kw, c in total_counts.items():
        flag = "" if c > 0 else "  <-- hiç bulunamadı, yazımı/büyük-küçük harfi kontrol edin"
        print(f"  {kw}: {c}{flag}")

    if all_leak_warnings:
        print("\n[!!! KRİTİK UYARI !!!] Aşağıdaki dosyalarda orijinal anahtar kelimelerden "
              "biri veya birkaçı hâlâ dosyanın içinde (görünmüyor olsa bile) tespit edildi. "
              "Bu dosyaları KULLANMAYIN, sebebini araştırın:")
        for w in all_leak_warnings:
            print("  " + w)
    elif not DRY_RUN:
        print("\n[tamam] Tüm çıktı dosyaları doğrulandı: orijinal anahtar kelimelerden hiçbiri dosyalarda kalmadı.")


if __name__ == "__main__":
    main()
